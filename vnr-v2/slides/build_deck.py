"""Builds VNR202_CaNgua_Showcase_v2.pptx: 12 slides following 05_Showcase_Slides_Script (10-12 min)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image
import os

HERE = os.path.dirname(os.path.abspath(__file__))
PAPER, INK, GREEN, GOLD, MUTED, RED = (RGBColor(*c) for c in [(0xF7, 0xF4, 0xEC), (0x24, 0x3D, 0x34), (0x28, 0x5E, 0x4D), (0xBF, 0x8B, 0x58), (0x66, 0x77, 0x6E), (0xA8, 0x43, 0x3B)])
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def bg(s, color=PAPER):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = color

def text(s, x, y, w, h, t, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font='Calibri'):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; r = p.add_run(); r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, font
        p.space_after = Pt(6)
    return tb

def bullets(s, x, y, w, h, items, size=18):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        head, _, rest = it.partition('::')
        r = p.add_run(); r.text = '▪ ' + head; r.font.size = Pt(size); r.font.bold = bool(rest); r.font.color.rgb = INK; r.font.name = 'Calibri'
        if rest:
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = INK; r2.font.name = 'Calibri'
        p.space_after = Pt(9)

def box(s, x, y, w, h, fill=GREEN):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background(); r.adjustments[0] = 0.08
    return r

def slide(title, section, n, notes, kicker=None):
    s = prs.slides.add_slide(blank); bg(s)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = GREEN; bar.line.fill.background()
    text(s, 0.6, 0.3, 9, 0.3, section.upper(), 12, True, GREEN)
    text(s, 0.6, 0.6, 12, 0.9, title, 32, True, INK, font='Cambria')
    text(s, 0.6, 7.05, 9, 0.3, 'VNR202 · Cá Ngựa Cách Mạng v2 · Fall 2026', 10, False, MUTED)
    text(s, 12.0, 7.05, 0.8, 0.3, str(n), 10, False, MUTED, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s

def picture(s, name, x, y, w=None, h=None, crop=None):
    path = os.path.join(HERE, 'img', name)
    if crop:
        im = Image.open(path).crop(crop); path = os.path.join(HERE, 'img', '_crop_' + name); im.save(path)
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

def card(s, x, y, w, h, title, body, fill=GREEN, tsize=16, bsize=13):
    box(s, x, y, w, h, fill)
    text(s, x + 0.2, y + 0.12, w - 0.4, 0.5, title, tsize, True, RGBColor(255, 255, 255))
    text(s, x + 0.2, y + 0.62, w - 0.4, h - 0.7, body, bsize, False, RGBColor(255, 255, 255))

# 1 ------------------------------------------------------------------------
s = prs.slides.add_slide(blank); bg(s, GREEN)
text(s, 0.8, 0.8, 11, 0.4, 'TRƯỜNG ĐẠI HỌC FPT · LÝ LUẬN CHÍNH TRỊ · VNR202 · FALL 2026', 13, True, RGBColor(0xCF, 0xE3, 0xD9))
text(s, 0.8, 1.9, 11.5, 1.6, 'Cá Ngựa Cách Mạng', 60, True, RGBColor(255, 255, 255), font='Cambria')
text(s, 0.8, 3.3, 11.5, 1.2, 'Trò chơi 6 đội ôn tập: Đảng Cộng sản Việt Nam ra đời và Cương lĩnh chính trị đầu tiên (2/1930)', 24, False, RGBColor(0xF7, 0xE7, 0xCF))
text(s, 0.8, 5.2, 11.5, 1.4, ['Nhóm: [TÊN NHÓM] · Lớp: [LỚP] · Topic ID: [MÃ]', 'Thành viên: [1] · [2] · [3] · [4] · [Nhóm trưởng]'], 16, False, RGBColor(255, 255, 255))
s.notes_slide.notes_text_frame.text = 'PHÚT 00-02 · Thành viên 1. Chào hội trường, nêu tên sản phẩm. Mở bằng câu hỏi: "Vì sao năm 1930 được gọi là sự lựa chọn của chính lịch sử?" Không đọc slide.'

# 2 ------------------------------------------------------------------------
s = slide('Vì sao cần một cách học khác về sự ra đời của Đảng?', 'Bối cảnh · Câu hỏi nghiên cứu trung tâm', 2,
          'PHÚT 00-02 · Thành viên 1. Nêu vấn đề: nội dung chương dày mốc thời gian, tên tổ chức, tên văn kiện nên sinh viên dễ học vẹt. Chốt Central Question và mục tiêu học tập.')
card(s, 0.6, 1.8, 5.9, 2.3, 'Vấn đề thực tế', 'Chương I có hàng chục mốc, 3 tổ chức cộng sản, 4 văn kiện hội nghị 1930. Sinh viên thuộc lòng rồi quên, khó thấy logic: vì sao các phong trào trước thất bại và vì sao chỉ con đường vô sản thành công.', GREEN)
card(s, 6.8, 1.8, 5.9, 2.3, 'Giải pháp: học bằng cuộc đua', 'Cá ngựa 6 người chơi. Mỗi lượt đổ xúc xắc và trả lời một câu hỏi từ giáo trình. Đúng thì nhận thẻ kỹ năng để tăng tốc hoặc chặn đối thủ. Về đích trước thì thắng.', GOLD)
box(s, 0.6, 4.5, 12.1, 1.9, RGBColor(0xE8, 0xF0, 0xE8))
text(s, 0.9, 4.65, 11.5, 0.4, 'CÂU HỎI NGHIÊN CỨU TRUNG TÂM', 12, True, GREEN)
text(s, 0.9, 5.05, 11.5, 1.3, '“Làm thế nào để sinh viên không chỉ nhớ mà còn hiểu vì sao sự ra đời của Đảng (2/1930) và Cương lĩnh chính trị đầu tiên là bước ngoặt tất yếu của cách mạng Việt Nam?”', 20, True, INK, font='Cambria')

# 3 ------------------------------------------------------------------------
s = slide('Căn cứ 1: Vì sao cách mạng Việt Nam cần một Đảng?', 'Cơ sở lý luận · Giáo trình Chương I, mục 1 (tr.36-49)', 3,
          'PHÚT 02-05 · Thành viên 2. Đi theo logic: thế giới → Việt Nam → các phong trào thất bại → nhu cầu một tổ chức tiên phong. Mỗi ý gắn 1 con số/ mốc, chỉ rõ thẻ/ câu hỏi trong game tương ứng. Chèn số Tập và Trang chuẩn vào Source Verification Log.')
bullets(s, 0.6, 1.6, 6.2, 5, [
    'Thế giới:: CNTB chuyển sang đế quốc chủ nghĩa; Cách mạng Tháng Mười Nga (1917); Quốc tế Cộng sản thành lập 3/1919; Đại hội II (1920) thông qua Luận cương dân tộc và thuộc địa của Lênin.',
    'Việt Nam:: Pháp tấn công Đà Nẵng 1/9/1858; Hiệp ước Patơnốt 6/6/1884; Liên bang Đông Dương 17/10/1887; hai cuộc khai thác thuộc địa (1897-1914, 1919-1929).',
    'Phân hóa xã hội:: nông dân >90% dân số; công nhân hơn 22 vạn (cuối 1929); mâu thuẫn chủ yếu: dân tộc Việt Nam >< thực dân Pháp và phong kiến phản động.'], 15)
box(s, 7.1, 1.6, 5.6, 4.9, RGBColor(0xFF, 0xFD, 0xF8))
text(s, 7.3, 1.7, 5.2, 0.4, 'Các phong trào yêu nước đều thất bại', 17, True, GREEN)
bullets(s, 7.3, 2.2, 5.2, 4.2, ['Cần Vương (1885-1896):: ngọn cờ phong kiến', 'Yên Thế (Hoàng Hoa Thám):: nặng “cốt cách phong kiến”', 'Đông Du, Duy Tân:: bạo động / cải cách dựa vào Pháp', 'Việt Nam Quốc dân Đảng, Yên Bái 2/1930:: manh động, ám sát cá nhân', 'Nguyên nhân chung:: thiếu đường lối đúng và tổ chức vững mạnh'], 14)

# 4 ------------------------------------------------------------------------
s = slide('Căn cứ 2: Nguyễn Ái Quốc chuẩn bị thành lập Đảng', 'Cơ sở lý luận · Giáo trình Chương I, mục 2 (tr.49-58)', 4,
          'PHÚT 02-05 · Thành viên 2. Dùng trục thời gian. Nhấn mạnh 3 lớp chuẩn bị: tư tưởng, chính trị, tổ chức. Liên hệ: thẻ "Đường Cách Mệnh" và "Chủ Nghĩa Làm Cốt" trong game lấy từ nội dung này.')
events = [('1911', 'Ra đi tìm\nđường cứu nước'), ('6/1919', 'Yêu sách 8 điểm\ntại Versailles'), ('7/1920', 'Đọc Luận cương\ncủa Lênin'), ('12/1920', 'Đại hội Tours:\ntán thành QT III'), ('6/1925', 'Hội VN Cách mạng\nThanh niên'), ('1927', 'Đường\ncách mệnh')]
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.55), Inches(11.9), Inches(0.06)); line.fill.solid(); line.fill.fore_color.rgb = GOLD; line.line.fill.background()
for i, (y, t) in enumerate(events):
    x = 0.7 + i * 2.05
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.45), Inches(2.4), Inches(0.36), Inches(0.36)); dot.fill.solid(); dot.fill.fore_color.rgb = GREEN; dot.line.fill.background()
    text(s, x, 1.85, 1.3, 0.4, y, 18, True, GREEN, PP_ALIGN.CENTER)
    text(s, x - 0.25, 2.95, 1.8, 0.9, t.split('\n'), 13, False, INK, PP_ALIGN.CENTER)
card(s, 0.6, 4.2, 3.9, 2.4, 'Tư tưởng', 'Báo Le Paria, Thanh niên; tố cáo chủ nghĩa thực dân; truyền bá chủ nghĩa Mác - Lênin vào phong trào công nhân và yêu nước.', GREEN)
card(s, 4.7, 4.2, 3.9, 2.4, 'Chính trị', 'Cách mạng giải phóng dân tộc là một bộ phận của cách mạng vô sản thế giới; công - nông là gốc của cách mệnh; cần có Đảng cách mệnh.', GOLD)
card(s, 8.8, 4.2, 3.9, 2.4, 'Tổ chức', 'Cộng sản đoàn (2/1925) → Hội VNCMTN (6/1925), 5 cấp tổ chức, hơn 10 lớp huấn luyện, báo Thanh niên (202 số), phong trào “Vô sản hóa”.', GREEN)

# 5 ------------------------------------------------------------------------
s = slide('Căn cứ 3: Hội nghị 2/1930 và Cương lĩnh đầu tiên', 'Cơ sở lý luận · Giáo trình Chương I, mục 3-4 (tr.58-70)', 5,
          'PHÚT 02-05 · Thành viên 2. Kể 3 bước: 1929 ba tổ chức → hội nghị 6/1-7/2/1930 tại Cửu Long (Hồng Kông) → Cương lĩnh. Đọc nguyên văn một câu trích. Chuyển sang slide 6: "Để học nội dung này, chúng em thiết kế một cuộc đua".')
bullets(s, 0.6, 1.6, 5.9, 5, [
    '1929:: Đông Dương Cộng sản Đảng (17/6), An Nam Cộng sản Đảng (11/1929), Đông Dương Cộng sản Liên đoàn (9/1929) → phân tán, cần hợp nhất.',
    'Hội nghị 6/1-7/2/1930:: Cửu Long (Hồng Kông), Nguyễn Ái Quốc chủ trì với tư cách phái viên Quốc tế Cộng sản; 4 đại biểu của 2 tổ chức; định tên Đảng Cộng sản Việt Nam.',
    'Văn kiện:: Chánh cương vắn tắt, Sách lược vắn tắt, Chương trình tóm tắt, Điều lệ vắn tắt.',
    'Ngày 3/2:: được Đại hội III (9/1960) chọn làm ngày kỷ niệm thành lập Đảng.'], 15)
card(s, 6.8, 1.6, 5.9, 1.55, 'Mục tiêu chiến lược', '“Tư sản dân quyền c.m và thổ địa c.m để đi tới xã hội cộng sản”', GREEN, 15, 14)
card(s, 6.8, 3.3, 2.85, 1.65, 'Nhiệm vụ', 'Chống đế quốc (hàng đầu) và chống phong kiến', GOLD, 15, 13)
card(s, 9.85, 3.3, 2.85, 1.65, 'Lực lượng', 'Công - nông là cơ bản, giai cấp công nhân lãnh đạo', GOLD, 15, 13)
card(s, 6.8, 5.1, 2.85, 1.5, 'Phương pháp', 'Bạo lực cách mạng của quần chúng', GREEN, 15, 13)
card(s, 9.85, 5.1, 2.85, 1.5, 'Ý nghĩa', 'Chấm dứt khủng hoảng đường lối; bước ngoặt vô cùng quan trọng', GREEN, 15, 13)

# 6 ------------------------------------------------------------------------
s = slide('Sản phẩm: Cá Ngựa Cách Mạng · luật chơi trong 30 giây', 'Sản phẩm sáng tạo', 6,
          'PHÚT 05-08 · Thành viên 3. Giải thích nhanh luật. Nhấn: đúng = đi + nhận thẻ; sai = vẫn đi nhưng không có thẻ, nên kiến thức quyết định công cụ để chặn đối thủ, không chỉ may rủi xúc xắc.')
steps = [('1', 'Đổ xúc xắc', 'Đội đến lượt đổ 1-6'), ('2', 'Trả lời', '1 câu 4 đáp án, 30 giây, từ giáo trình'), ('3', 'Đúng', 'Đi đủ ô + nhận 1 thẻ kỹ năng'), ('4', 'Sai / hết giờ', 'Vẫn đi tiếp, không có thẻ'), ('5', 'Về đích', 'Ai tới ĐÍCH 3/2/1930 trước thắng')]
for i, (n, t, b) in enumerate(steps):
    x = 0.6 + i * 2.45
    box(s, x, 1.7, 2.25, 2.7, GREEN if i != 3 else RED)
    text(s, x + 0.15, 1.8, 2, 0.8, n, 36, True, RGBColor(0xF7, 0xE7, 0xCF), font='Cambria')
    text(s, x + 0.15, 2.65, 2, 0.5, t, 17, True, RGBColor(255, 255, 255))
    text(s, x + 0.15, 3.2, 1.95, 1.2, b, 13, False, RGBColor(255, 255, 255))
bullets(s, 0.6, 4.8, 12, 2, ['Bàn cờ:: vòng đua chung 24 ô + đường về nhà 4 ô an toàn; đáp xuống ô có ngựa đội khác (trừ ô ★) thì húc lùi 3 ô.', 'Mốc lịch sử:: mỗi chặng ngựa đi qua bật một mốc (1858, 1911, 1920, 1925, 1929).', 'Quy mô:: 2-6 người chơi, mỗi người một ngựa trên một laptop; màn chiếu riêng cho lớp; 20-30 người cùng xem.'], 15)

# 7 ------------------------------------------------------------------------
s = slide('DEMO trực tiếp: một lượt chơi', 'Sản phẩm sáng tạo · Demo', 7,
          'PHÚT 05-08 · Thành viên 3. DEMO THẬT: (1) mở màn chiếu /rooms/<mã>/board; (2) một đội đổ xúc xắc; (3) đọc câu hỏi, trả lời; (4) chỉ lời giải thích và số trang hiện sau mỗi lượt; (5) nhấn vào bảng xếp hạng. Chuẩn bị phương án dự phòng: video quay màn hình 60 giây.')
picture(s, '06-projector.png', 0.6, 1.6, w=6.9)
picture(s, '02-lobby-host.png', 7.7, 1.6, w=5.1)
text(s, 0.6, 5.75, 7.4, 1.2, ['Màn chiếu cho cả lớp: bàn cờ, lượt, câu hỏi, bảng xếp hạng, diễn biến.', 'Phòng chờ: mỗi người chọn 1 trong 6 ngựa. Trên laptop của mình, người chơi thấy câu hỏi, 4 đáp án và thẻ kỹ năng riêng (đối thủ chỉ thấy số thẻ).'], 14, False, INK)

# 8 ------------------------------------------------------------------------
s = slide('10 thẻ kỹ năng: dùng lên đối thủ để chặn đường về đích', 'Sản phẩm sáng tạo · Gắn với nội dung giáo trình', 8,
          'PHÚT 05-08 · Thành viên 3. Chọn 3 thẻ để minh họa: Chia Để Trị (Pháp chia ba kỳ), Mua Chuộc (Sarraut), Chủ Nghĩa Làm Cốt (Đường cách mệnh). Thông điệp: mỗi thẻ là một mảnh kiến thức, chơi game là nhắc lại lịch sử.')
skills = [('✂️ Chia Để Trị ×3', 'Đối thủ mất lượt kế tiếp', 'Pháp chia ba kỳ'), ('🚧 Thuế Khóa Nặng Nề ×3', 'Đối thủ đi tối đa 2 bước', 'Thuế khóa, bóc lột'), ('⏪ Đàn Áp Khởi Nghĩa ×3', 'Đối thủ lùi 4 ô', 'Yên Bái bị dập tắt'), ('🔄 Hoán Đổi Vận Mệnh ×2', 'Đổi chỗ với ngựa khác', 'Đổi thay khi có đường lối đúng'), ('🕵️ Mua Chuộc ×2', 'Lấy 1 thẻ của đối thủ', 'Sarraut mua chuộc Nguyễn Ái Quốc'),
          ('🚩 Đường Cách Mệnh ×3', 'Lần đi sau +3 bước', 'Tác phẩm 1927'), ('🎲 Nhãn Quan Sắc Bén ×2', 'Đổ 2 xúc xắc lấy số cao', 'Nhãn quan chính trị'), ('🛡️ Chủ Nghĩa Làm Cốt ×3', 'Khiên chặn 1 kỹ năng', '“Đảng muốn vững thì phải có chủ nghĩa làm cốt”'), ('💡 Soi Sáng ×3', 'Mở gợi ý khi trả lời', 'Chủ nghĩa Mác - Lênin soi đường'), ('🎯 Thu Hẹp Vòng Vây ×2', 'Loại 2 đáp án sai', 'Loại các con đường thất bại')]
for i, (n, e, t) in enumerate(skills):
    col, row = i % 5, i // 5
    x, y = 0.6 + col * 2.45, 1.7 + row * 2.7
    box(s, x, y, 2.3, 2.5, RED if i < 5 else GREEN)
    text(s, x + 0.12, y + 0.1, 2.1, 0.7, n, 14, True, RGBColor(255, 255, 255))
    text(s, x + 0.12, y + 0.85, 2.1, 0.7, e, 13, True, RGBColor(0xF7, 0xE7, 0xCF))
    text(s, x + 0.12, y + 1.5, 2.1, 0.95, t, 11, False, RGBColor(255, 255, 255))
text(s, 0.6, 6.7, 12, 0.3, 'Đỏ: chặn đối thủ · Xanh: tự thân / hỗ trợ khi trả lời · Tay tối đa 5 thẻ, mỗi lượt dùng 1 thẻ trước khi đổ xúc xắc.', 12, False, MUTED)

# 9 ------------------------------------------------------------------------
s = slide('Thử nghiệm thực tế và phản hồi người chơi', 'Tương tác · Phụ lục 09_Interaction_Feedback_Evidence', 9,
          'PHÚT 08-10 · Thành viên 4. BẮT BUỘC thay các ô [..] bằng số liệu thật từ buổi chơi thử (tối thiểu 15-20 người) và đính kèm form khảo sát, ảnh chụp. Không để trống hoặc bịa số. Nêu rõ 2-3 phản hồi đã dẫn tới thay đổi ở v2.')
tiles = [('[SỐ]', 'người tham gia thử nghiệm (mục tiêu ≥ 15-20)'), ('[%]', 'tự đánh giá nhớ mốc 1930 tốt hơn sau khi chơi'), ('[x/5]', 'điểm hài lòng trung bình'), ('[%]', 'trả lời đúng ở lượt đầu và lượt cuối (so sánh)')]
for i, (v, l) in enumerate(tiles):
    x = 0.6 + i * 3.05
    box(s, x, 1.7, 2.85, 2.1, RGBColor(0xE8, 0xF0, 0xE8))
    text(s, x + 0.15, 1.8, 2.6, 0.9, v, 36, True, GREEN, font='Cambria')
    text(s, x + 0.15, 2.7, 2.6, 1, l, 13, False, INK)
text(s, 0.6, 4.1, 12, 0.4, 'Phản hồi tiêu biểu → thay đổi ở v2 (điền từ form thật)', 17, True, GREEN)
bullets(s, 0.6, 4.6, 12, 2.2, ['“[Trích phản hồi 1]”:: → [Thay đổi đã làm]', '“[Trích phản hồi 2]”:: → [Thay đổi đã làm]', '“[Trích phản hồi 3]”:: → [Thay đổi đã làm]'], 15)

# 10 -----------------------------------------------------------------------
s = slide('Nhật ký cải tiến sản phẩm: từ v1 đến v2', 'Quy trình thiết kế · 04_Product_Prototype_DevelopmentLog', 10,
          'PHÚT 08-10 · Thành viên 4. So sánh v1 và v2. Giải thích lý do chuyển từ đấu máu 2 đội sang cá ngựa 6 đội: lớp đông, cần cả lớp cùng tham gia và có tương tác chặn nhau. Liệt kê rõ phần AI hỗ trợ trong 07_AI_Declaration.')
card(s, 0.6, 1.7, 5.9, 4.4, 'v1 · Đấu Trường Tri Thức', ['2 đội đấu trực tiếp, thanh máu (HP)', 'Thẻ câu hỏi, thẻ đáp án, 10 kỹ năng công / thủ', 'Bán kết + chung kết, chỉ 4 đội', 'Nội dung: Tư tưởng Hồ Chí Minh', 'Hạn chế: nhiều đội ngồi chờ lâu; khó cho lớp 20-30 người'], GOLD)
card(s, 6.8, 1.7, 5.9, 4.4, 'v2 · Cá Ngựa Cách Mạng', ['6 người đua cùng lúc trên một bàn cờ', 'Đúng → đi + nhận thẻ; sai → đi nhưng không thẻ', '10 thẻ kỹ năng: chặn đối thủ, tăng tốc, khiên, hỗ trợ', '40 câu hỏi lấy từ Chương I, có giải thích + số trang', 'Backend NestJS, màn chiếu riêng, đo thử 30 client'], GREEN)
text(s, 0.6, 6.3, 12, 0.5, 'Giới hạn hiện tại: dữ liệu phòng lưu trong bộ nhớ (mất khi khởi động lại); hướng phát triển: Redis, WebSocket, thêm bộ câu hỏi các chương khác.', 13, False, MUTED)

# 11 -----------------------------------------------------------------------
s = slide('Thông điệp cốt lõi và bài học rút ra', 'Tổng kết', 11,
          'PHÚT 10-12 · Nhóm trưởng. Đóng vòng: trả lời lại Central Question bằng 3 ý. Mỗi thành viên nêu 1 bài học cá nhân (liên hệ rubric 2.3). Không quá 2 phút.')
card(s, 0.6, 1.7, 3.9, 3.3, 'Lý luận', 'Đảng ra đời (2/1930) là kết quả của sự kết hợp chủ nghĩa Mác - Lênin với phong trào công nhân và phong trào yêu nước; Cương lĩnh đầu tiên chọn con đường cách mạng vô sản.', GREEN, 17, 14)
card(s, 4.7, 1.7, 3.9, 3.3, 'Sản phẩm', 'Biến kiến thức thành lựa chọn có hậu quả: mỗi đáp án đúng trao công cụ chặn đối thủ; lời giải thích kèm trang giáo trình sau mỗi lượt.', GOLD, 17, 14)
card(s, 8.8, 1.7, 3.9, 3.3, 'Bài học', '[Mỗi thành viên 1 câu]\n[Khó khăn kỹ thuật/nội dung và cách giải quyết]\n[Điều sẽ làm ở v3]', GREEN, 17, 14)
box(s, 0.6, 5.4, 12.1, 1.2, RGBColor(0xE8, 0xF0, 0xE8))
text(s, 0.9, 5.55, 11.5, 0.9, '“Việc thành lập Đảng là một bước ngoặt vô cùng quan trọng trong lịch sử cách mạng Việt Nam ta.” (Hồ Chí Minh)', 20, True, INK, font='Cambria')

# 12 -----------------------------------------------------------------------
s = slide('Sẵn sàng phản biện 3C: Context · Concept · Conflict', 'Q&A', 12,
          'PHÚT 10-12 · Nhóm trưởng. Cam kết No-AI live rebuttal: trả lời bằng hiểu biết của nhóm. Phân công trước: ai trả lời Context, ai Concept, ai Conflict. Chuẩn bị đáp án 3-4 câu mỗi nhóm; dùng phần giải thích và trang trong ngân hàng câu hỏi.')
card(s, 0.6, 1.7, 3.9, 4.7, 'Context (Bối cảnh)', ['Vì sao đầu thế kỷ XX các phong trào yêu nước đều thất bại?', 'Vì sao Nguyễn Ái Quốc chọn con đường Lênin?', 'Vì sao năm 1929 cần hợp nhất ba tổ chức?'], GREEN, 17, 14)
card(s, 4.7, 1.7, 3.9, 4.7, 'Concept (Khái niệm)', ['Cương lĩnh đầu tiên gồm những văn kiện nào, “vắn tắt” nghĩa là gì?', 'Tư sản dân quyền cách mạng và thổ địa cách mạng là gì?', 'Vì sao chống đế quốc đặt hàng đầu?'], GOLD, 17, 14)
card(s, 8.8, 1.7, 3.9, 4.7, 'Conflict (Mâu thuẫn)', ['Trò chơi có làm nhẹ tính nghiêm túc của lịch sử không? Nhóm xử lý thế nào?', 'Đáp án sai mà vẫn được đi: có khuyến khích học vẹt không?', 'Thẻ “Mua Chuộc”, “Đàn Áp” gắn với lịch sử có phù hợp?'], RED, 17, 14)
text(s, 0.6, 6.55, 12, 0.4, 'Cảm ơn hội trường · Demo và mã nguồn: [LINK DRIVE / REPO]', 14, True, GREEN)

out = os.path.join(HERE, 'VNR202_CaNgua_Showcase_v2.pptx'); prs.save(out); print(out)
